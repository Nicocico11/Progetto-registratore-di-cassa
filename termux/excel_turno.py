import json, os
from datetime import datetime, date, time, timedelta

# Compila il modello Excel del distributore ("TURNO NUOVO") con i dati del turno appena chiuso
# e prepara ("imbastisce") il file del turno successivo.
# Si scrive SOLO nelle caselle bianche: formule e protezione del foglio restano intatte.

MODELLO = os.path.expanduser('~/.termux/tasker/modello_turno.xlsx')
PATH_STATO = os.path.expanduser('~/stato_cassa.json')  # contatore AdBlue, taniche, avanzo tra un turno e l'altro

# Caselle del modello (nell'ordine in cui le sommano le formule del foglio)
DANEA = [(f'E{r}', f'H{r}') for r in range(2, 30)]                                  # descrizione, importo
TELEFAX = [f'{c}{r}' for r in range(21, 26) for c in 'ABCD']
ADBLUE_LITRI = [f'{c}17' for c in 'IJKLMNO'] + [f'{c}18' for c in 'IJKLMN']
SCONTRINI_POS = [f'{c}{r}' for r in range(27, 33) for c in 'STUV']
CREDITI_CLIENTI = [(f'I{r}', f'L{r}') for r in range(8, 16)]   # nome, importo
CREDITI_RISCOSSI = [(f'M{r}', f'O{r}') for r in range(8, 16)]
# La notte porta la data del giorno in cui finisce: pomeriggio del 4 -> notte del 5 -> mattina del 5
SUCCESSIVO = {'Mattina': ('Pomeriggio', 0), 'Pomeriggio': ('Notte', 1), 'Notte': ('Mattina', 0)}


def turno_di_prova():
    """Vero se il turno aperto è di prova ("apertura turno prova")."""
    try:
        with open(os.path.expanduser('~/turno_corrente.json'), encoding='utf-8') as f:
            return bool(json.load(f).get('prova'))
    except Exception:
        return False


def path_stato():
    # Turno di prova: contatore, taniche e orari in un file a parte, quello vero non si tocca
    return PATH_STATO.replace('.json', '_prova.json') if turno_di_prova() else PATH_STATO


def leggi_stato():
    try:
        with open(path_stato(), encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def salva_stato(stato):
    with open(path_stato(), 'w', encoding='utf-8') as f:
        json.dump(stato, f)


def riempi(ws, caselle, valori, avvisi, nome, nomi=None):
    """Scrive i valori nelle caselle, uno per casella. Se le caselle finiscono si riparte dalla
    prima sommando (la 21ª voce si somma alla 1ª, la 22ª alla 2ª...): il totale resta giusto.
    nomi: per i crediti, caselle del nome accanto all'importo (i nomi si uniscono con " + ")."""
    coppie = [(round(v, 2), n) for v, n in zip(valori, nomi or [None] * len(valori)) if v]
    if len(coppie) > len(caselle):
        avvisi.append(f"{nome}: {len(coppie)} voci per {len(caselle)} caselle: "
                      "dalla prima casella in poi alcune sono la somma di due voci")
    somme, testi = {}, {}
    for i, (v, n) in enumerate(coppie):
        k = i % len(caselle)
        somme[k] = round(somme.get(k, 0.0) + v, 2)
        if n is not None:
            testi[k] = f"{testi[k]} + {n}" if k in testi else n
    for k, v in somme.items():
        if nomi is None:
            ws[caselle[k]] = v
        else:
            ws[caselle[k][0]] = testi[k]
            ws[caselle[k][1]] = v


def nome_file(giorno, tipo, prova=False):
    return f"{giorno.strftime('%d_%m_%Y')}_{tipo.lower()}{'_TEST' if prova else ''}.xlsx"


def ora_da_testo(testo):
    try:
        h, m, s = (int(x) for x in str(testo).split()[0].split(':'))
        return time(h, m, s)
    except Exception:
        return None


def crea_excel(righe, turno, orario_terminale, cartella):
    """Crea <data>_<turno>.xlsx (turno concluso) e il file del turno successivo.
    Restituisce (lista di messaggi, stato aggiornato).
    "ORA CHIUSURA" (I24) è l'orario del terminale pompe del turno PRECEDENTE:
    quello inserito oggi va nel file del turno dopo."""
    import openpyxl
    avvisi = []
    stato = leggi_stato()
    voci = []
    for r in righe:
        try:
            voci.append(json.loads(r[1]))
        except Exception:
            pass
    imp = lambda v: float(v.get('importo', 0) or 0)
    q = lambda v: float(v.get('quantita', 1) or 1)
    giorno = date.fromisoformat(turno['data_file'])

    wb = openpyxl.load_workbook(MODELLO)
    wb.calculation.fullCalcOnLoad = True  # i totali del foglio li calcola Excel all'apertura
    ws = wb.active

    # Intestazione: data, turno, ora chiusura del turno precedente
    ws['I20'] = datetime.combine(giorno, time())
    ws['I22'] = turno['tipo'].upper()
    if ora_da_testo(stato.get('orario_chiusura')):
        ws['I24'] = ora_da_testo(stato['orario_chiusura'])
    else:
        avvisi.append("ora chiusura del turno precedente sconosciuta: scrivila a mano")
    ora_oggi = ora_da_testo(orario_terminale)
    if not ora_oggi:
        avvisi.append("orario terminale non inserito: nel file del turno dopo va scritto a mano")

    # DANEA: prodotti market raggruppati ("3x RED BULL")
    gruppi = {}
    for v in voci:
        if v.get('reparto') == 'Market':
            n, tot = gruppi.get(v['categoria'], (0.0, 0.0))
            gruppi[v['categoria']] = (n + q(v), tot + imp(v))
    danea = sorted(gruppi.items(), key=lambda x: -x[1][1])
    if len(danea) > len(DANEA):
        avvisi.append(f"DANEA: {len(danea)} prodotti, righe {len(DANEA)}: gli ultimi sono sommati in 'ALTRI'")
        resto = danea[len(DANEA) - 1:]
        danea = danea[:len(DANEA) - 1] + [('ALTRI', (sum(x[1][0] for x in resto), sum(x[1][1] for x in resto)))]
    for (c_desc, c_imp), (nome, (n, tot)) in zip(DANEA, danea):
        ws[c_desc] = nome if n == 1 else f"{n:g}x {nome}"
        ws[c_imp] = round(tot, 2)

    # TELEFAX, litri AdBlue sfuso
    riempi(ws, TELEFAX, [imp(v) for v in voci if v.get('reparto') == 'Fax'], avvisi, "TELEFAX")
    litri = [q(v) for v in voci if v.get('reparto') == 'AdBlue' and v.get('unita') == 'l']
    riempi(ws, ADBLUE_LITRI, litri, avvisi, "ADBLUE")

    # Crediti clienti (non pagati ora) e crediti riscossi (vecchi crediti pagati oggi)
    for reparto, caselle, nome in (('Credito cliente', CREDITI_CLIENTI, "CREDITI CLIENTI"),
                                   ('Credito riscosso', CREDITI_RISCOSSI, "CREDITI RISCOSSI")):
        crediti = [v for v in voci if v.get('reparto') == reparto]
        riempi(ws, caselle, [imp(v) for v in crediti], avvisi, nome,
               nomi=[v.get('cliente') or '?' for v in crediti])

    # Abbuoni (centesimi in meno) e resti lasciati dai clienti (centesimi in più): non hanno caselle,
    # finiscono nella DIFFERENZA del foglio. Una nota (A39) spiega da dove viene.
    note = []
    for reparto, titolo in (('Sconto', 'ABBUONI'), ('Resto lasciato', 'RESTI LASCIATI DAI CLIENTI')):
        lista = [imp(v) for v in voci if v.get('reparto') == reparto]
        if lista:
            totale = round(sum(lista), 2)
            note.append(f"{titolo}: {'+' if totale > 0 else '-'} {abs(totale):.2f} € ({len(lista)} {'volta' if len(lista) == 1 else 'volte'})".replace('.', ','))
    if note:
        ws['A39'] = " | ".join(note) + " - compaiono nella differenza"

    # Pagamenti con carta: i tre POS "esterni" come totali, il POS della cassa uno scontrino per vendita
    metodo = lambda v: {'Carta carburante': 'Petrolifere'}.get(v.get('metodo_pagamento'), v.get('metodo_pagamento'))
    for casella, nome in (('D9', 'Petrolifere'), ('D12', 'POS nero'), ('D14', 'POS bianco')):
        totale = round(sum(imp(v) for v in voci if metodo(v) == nome), 2)
        if totale:
            ws[casella] = totale
    per_scontrino = {}
    for v in voci:
        if metodo(v) == 'POS cassa':
            chiave = v.get('transazione') or id(v)
            per_scontrino[chiave] = per_scontrino.get(chiave, 0.0) + imp(v)
    riempi(ws, SCONTRINI_POS, list(per_scontrino.values()), avvisi, "SCONTRINI POS")

    # Contatore AdBlue e taniche: iniziali dal turno prima, finali calcolati dalle vendite
    contatore_finale = taniche_attuali = None
    if stato.get('contatore') is not None:
        ws['O20'] = stato['contatore']
        contatore_finale = round(stato['contatore'] + sum(litri), 2)
        ws['O19'] = contatore_finale
    else:
        avvisi.append("contatore AdBlue iniziale sconosciuto: di' \"contatore adblue …\" all'apertura")
    vendute = sum(q(v) for v in voci if v.get('reparto') == 'AdBlue' and v.get('unita') != 'l')
    if stato.get('taniche') is not None:
        ws['K30'] = stato['taniche']
        taniche_attuali = stato['taniche'] - vendute
        ws['K31'] = taniche_attuali
    else:
        avvisi.append("taniche AdBlue sconosciute: di' \"contatore taniche …\" all'apertura")

    # Cassa: avanzo precedente; nei "spiccioli cassetto" i contanti che dovrebbero esserci
    avanzo = turno.get('avanzo')
    if avanzo is not None:
        ws['D7'] = avanzo
    # Cassaforte: va in "IN CASSAFORTE" (I28), il foglio la somma da solo all'avanzo attuale
    contanti = sum(imp(v) for v in voci if v.get('metodo_pagamento') == 'Contanti')
    totale_cassa = round((avanzo or 0) + contanti - (turno.get('versamento') or 0), 2)
    cassaforte = turno.get('cassaforte') or 0
    if cassaforte:
        ws['I28'] = cassaforte
    ws['D34'] = round(totale_cassa - cassaforte, 2)

    os.makedirs(cartella, exist_ok=True)
    path_turno = os.path.join(cartella, nome_file(giorno, turno['tipo'], turno.get('prova') or turno.get('nomi_test')))
    wb.save(path_turno)

    # Turno successivo "imbastito"
    tipo_dopo, giorni = SUCCESSIVO[turno['tipo']]
    giorno_dopo = giorno + timedelta(days=giorni)
    cassetto = turno.get('contati') if turno.get('contati') is not None else totale_cassa - cassaforte
    avanzo_dopo = round(cassetto + cassaforte, 2)
    wb2 = openpyxl.load_workbook(MODELLO)
    wb2.calculation.fullCalcOnLoad = True
    ws2 = wb2.active
    ws2['I20'] = datetime.combine(giorno_dopo, time())
    ws2['I22'] = tipo_dopo.upper()
    ws2['D7'] = avanzo_dopo
    if ora_oggi:
        ws2['I24'] = ora_oggi
    if cassaforte:
        ws2['I28'] = cassaforte
    if contatore_finale is not None:
        ws2['O20'] = contatore_finale
    if taniche_attuali is not None:
        ws2['K30'] = taniche_attuali
    path_dopo = os.path.join(cartella, nome_file(giorno_dopo, tipo_dopo, turno.get('prova') or turno.get('nomi_test')))
    if not os.path.exists(path_dopo):  # non sovrascrivere un turno già compilato
        wb2.save(path_dopo)

    stato.update({'contatore': contatore_finale if contatore_finale is not None else stato.get('contatore'),
                  'taniche': taniche_attuali if taniche_attuali is not None else stato.get('taniche'),
                  'avanzo': avanzo_dopo,
                  'orario_chiusura': ora_oggi.strftime('%H:%M:%S') if ora_oggi else None})
    salva_stato(stato)

    breve = lambda p: "Download/" + os.path.relpath(p, os.path.expanduser('~/storage/downloads'))
    messaggi = [f"📗 Excel: {breve(path_turno)}",
                f"📘 Turno dopo: {breve(path_dopo)}"]
    return messaggi + [f"⚠️ {a}" for a in avvisi], stato
